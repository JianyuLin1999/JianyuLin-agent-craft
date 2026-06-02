#!/usr/bin/env python3
"""docx 后处理 —— 三线表 + 图注左对齐 + front matter 分页 + 域自动更新。

pandoc 生成的 docx 还差几件事才达专业报告格式,这里逐项补:
  1. 三线表:每个表清掉竖线和内部横线,只留顶线(粗)/ 表头下线(细)/ 底线(粗)
  2. 图注 / 表注左对齐(reference style 没生效时兜底)
  3. front matter 分页:标题页独占 + 目录独占 + 摘要独占 + 正文连续流
  4. 域自动更新:设 updateFields,让 soffice / Word 打开时自动刷新目录(否则目录空白)

用法:python postprocess-docx.py <input.docx> [output.docx]
依赖:python-docx。
"""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

inp = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else inp
d = Document(inp)

def _edge(name, val, sz):
    e = OxmlElement(f'w:{name}')
    e.set(qn('w:val'), val)
    e.set(qn('w:sz'), str(sz))      # 1/8 pt:12=1.5pt 粗,6=0.75pt 细
    e.set(qn('w:space'), '0')
    e.set(qn('w:color'), '000000' if val == 'single' else 'auto')
    return e

def three_line(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    b.append(_edge('top', 'single', 12))        # 粗顶线
    b.append(_edge('bottom', 'single', 12))     # 粗底线
    for e in ('left', 'right', 'insideH', 'insideV'):
        b.append(_edge(e, 'none', 0))           # 无竖线、无内部横线
    tblPr.append(b)
    for cell in table.rows[0].cells:            # 表头行下细线
        tcPr = cell._tc.get_or_add_tcPr()
        oldb = tcPr.find(qn('w:tcBorders'))
        if oldb is not None: tcPr.remove(oldb)
        tcB = OxmlElement('w:tcBorders')
        tcB.append(_edge('bottom', 'single', 6))
        tcPr.append(tcB)

# 1. 三线表
for t in d.tables:
    three_line(t)

# 2. 图注 / 表注左对齐
for p in d.paragraphs:
    if 'caption' in p.style.name.lower():
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# 3. front matter 分页:标题页 / 目录 / 摘要 各独占,正文连续
#    目录标题(TOC Heading)段前分页 → 标题页独占
for p in d.paragraphs:
    if p.style.name == 'TOC Heading':
        p.paragraph_format.page_break_before = True
        break
#    前两个一级标题(摘要 / 第一主章节)段前分页 → 目录独占 + 摘要独占
h1 = [p for p in d.paragraphs if p.style.name == 'Heading 1']
for p in h1[:2]:
    p.paragraph_format.page_break_before = True

# 4. 域自动更新(目录):打开文档时刷新所有域,否则 soffice 转 PDF 目录空白
settings = d.settings.element
if settings.find(qn('w:updateFields')) is None:
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    settings.append(uf)

d.save(out)
print(f"✓ {out} — {len(d.tables)} 表三线表 · 图注左对齐 · front matter 分页 · 目录域自动更新")
